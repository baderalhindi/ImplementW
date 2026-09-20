#!/usr/bin/env python3
"""Build ptbc-tbc-tracker.xlsx from the CSV registers in this directory.

The CSVs are the tracked source of record; the workbook is a generated view for
PMO use. Edit the CSVs, then re-run this script.

    python3 docs/governance/build-tracker.py

Requires openpyxl.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
OUTPUT = HERE / "ptbc-tbc-tracker.xlsx"

SHEETS = [
    ("PTBC Themes", "ptbc-themes.csv"),
    ("Task Cross-Reference", "ptbc-task-crossreference.csv"),
    ("Source TBC Crosswalk", "source-tbc-crosswalk.csv"),
    ("Unassigned Gated Values", "unassigned-gated-values.csv"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MAX_WIDTH = 60


def add_sheet(workbook, title, csv_path):
    sheet = workbook.create_sheet(title)
    with csv_path.open(newline="") as handle:
        rows = list(csv.reader(handle))

    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for index, column in enumerate(zip(*rows), start=1):
        width = max(len(str(value)) for value in column)
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, MAX_WIDTH)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def main():
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, filename in SHEETS:
        add_sheet(workbook, title, HERE / filename)
    workbook.save(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
